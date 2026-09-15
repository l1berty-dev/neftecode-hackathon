import re
from datetime import UTC, datetime
from pathlib import Path

import openpyxl

from neftecode_hackathon.data.prepare import (
    _iter_telemetry_chunks,
    _value_quality,
    load_lims,
    load_pak,
    parse_point_header,
)

POINT_HEADER = "Установка 'Гидроочистка'.. Точка отбора '2'. Продукт 'Дизельное топливо'"


def test_point_header_preserves_installation_point_and_product() -> None:
    point = parse_point_header(POINT_HEADER)

    assert point.installation == "Гидроочистка"
    assert point.point == "2"
    assert point.product == "Дизельное топливо"
    assert point.point_id == "ht:sample_point_2"


def test_telemetry_is_sorted_namespaced_and_keeps_307(tmp_path: Path) -> None:
    source = tmp_path / "telemetry.csv"
    source.write_text(
        ",date,F1,T2\n1,2026-01-01 00:10:00,-1,307\n0,2026-01-01 00:00:00,2,10\n",
        encoding="utf-8",
    )

    chunks = list(
        _iter_telemetry_chunks(
            source,
            namespace="ht",
            descriptions={"F1": "Расход сырья", "T2": "Температура"},
            source_timezone="Europe/Moscow",
            delay_minutes=0,
            timestamp_column="date",
            timestamp_format="%Y-%m-%d %H:%M:%S",
            chunksize=1,
            service_pattern=re.compile(r"^Unnamed:|^$"),
            audit_exact_values=[307],
        )
    )

    telemetry = [row for frame, _, _ in chunks for row in frame.to_dict("records")]
    assert [row["measured_at"] for row in telemetry[::2]] == [
        datetime(2025, 12, 31, 21, 0, tzinfo=UTC),
        datetime(2025, 12, 31, 21, 10, tzinfo=UTC),
    ]
    assert {row["signal_id"] for row in telemetry} == {"ht:F1", "ht:T2"}
    kept_307 = next(row for row in telemetry if row["value"] == 307)
    assert kept_307["quality"] == "valid"
    negative_flow = next(row for row in telemetry if row["value"] == -1)
    assert negative_flow["quality"] == "suspect"
    assert negative_flow["issues"] == '["negative_flow"]'
    assert all(audit["service_columns_removed"] == ["Unnamed: 0"] for _, _, audit in chunks)


def test_quality_rules_are_effective_configuration() -> None:
    quality, issues = _value_quality(
        "recovery_350",
        350.0,
        (),
        nonnegative_analytes=(),
        percentage_analytes=(),
    )

    assert quality == "valid"
    assert issues == "[]"


def test_lims_pairs_are_independent_and_text_becomes_event(tmp_path: Path) -> None:
    path = tmp_path / "lims.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([POINT_HEADER, None, None, None])
    sheet.append(["Mg.Sulfur", None, "D15", None])
    sheet.append(["мг/кг", None, "°С", None])
    sheet.append(["Количество значений:", 2, "Количество значений:", 1])
    sheet.append([datetime(2026, 1, 1, 10), 5.0, None, None])
    sheet.append([datetime(2026, 1, 2, 10), "Pt Created", datetime(2026, 1, 3, 10), 840.0])
    workbook.save(path)

    analyses, events, dictionary, audit = load_lims(
        path, source_timezone="Europe/Moscow", delay_minutes=240
    )

    assert set(analyses["signal_id"]) == {
        "lab:ht.point_2.sulfur",
        "lab:ht.point_2.density_d15",
    }
    sulfur = analyses.loc[analyses["signal_id"] == "lab:ht.point_2.sulfur"].iloc[0]
    assert sulfur["measured_at"] == datetime(2026, 1, 1, 7, tzinfo=UTC)
    assert sulfur["available_at"] == datetime(2026, 1, 1, 11, tzinfo=UTC)
    assert events.iloc[0]["raw_value"] == "Pt Created"
    assert events.iloc[0]["event_type"] == "source_text"
    density_dictionary = next(
        item for item in dictionary if item["signal_id"] == "lab:ht.point_2.density_d15"
    )
    assert density_dictionary["verification_status"] == "organizer_confirmed_unit_correction"
    assert density_dictionary["canonical_unit"] == "kg/m³"
    density = analyses.loc[analyses["signal_id"] == "lab:ht.point_2.density_d15"].iloc[0]
    assert density["quality"] == "valid"
    assert density["original_unit"] == "°С"
    assert density["issues"] == '["source_unit_corrected_from_analyte"]'
    assert audit[0]["count_matches_declared"] is True


def test_pak_uses_two_pairs_separated_by_empty_column(tmp_path: Path) -> None:
    path = tmp_path / "pak.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["24-2000:Mg.Sulfur", None, None, "24-2000:D15", None])
    sheet.append(["ppm", None, None, "кг/м3", None])
    sheet.append([datetime(2026, 1, 1, 10), 5.0, None, datetime(2026, 1, 2, 10), 840.0])
    workbook.save(path)

    analyses, events, dictionary, audit = load_pak(
        path, source_timezone="Europe/Moscow", delay_minutes=0
    )

    assert events.empty
    assert set(analyses["signal_id"]) == {
        "pak:ht.product_sulfur",
        "pak:ht.product_density_d15",
    }
    assert set(analyses["unit"]) == {"mg/kg", "kg/m³"}
    assert dictionary[0]["conversion"] == "1 mass ppm = 1 mg/kg (assumption)"
    assert [item["numeric_count"] for item in audit] == [1, 1]


def test_pak_rejects_nonempty_separator_column(tmp_path: Path) -> None:
    path = tmp_path / "pak.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["24-2000:Mg.Sulfur", None, "unexpected", "24-2000:D15", None])
    workbook.save(path)

    try:
        load_pak(path, source_timezone="Europe/Moscow", delay_minutes=0)
    except ValueError as error:
        assert "separator column C" in str(error)
    else:
        raise AssertionError("nonempty PAK separator must be rejected")
