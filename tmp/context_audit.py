import collections
import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import pypdfium2 as pdfium

root = Path("context")
out = Path("tmp/context_audit")
out.mkdir(parents=True, exist_ok=True)
for p in root.glob("*.pdf"):
    doc = pdfium.PdfDocument(str(p))
    for i, page in enumerate(doc):
        page.render(scale=1.2 if "схемы" in p.name else 1.5).to_pil().save(
            out / f"{p.stem}-{i + 1}.png"
        )
for p in root.glob("*.xlsx"):
    w = openpyxl.load_workbook(p, read_only=True, data_only=True)
    if "Теги" in p.name:
        continue
    print("\nFILE", p.name, flush=True)
    for s in w:
        rows = list(s.values)
        print("SHEET", s.title, "DIM", s.max_row, s.max_column)
        group = None
        start = 2 if "ПАК" in p.name else 4
        for j in [0, 3] if "ПАК" in p.name else range(0, s.max_column - 1, 2):
            heads = [r[j : j + 2] for r in rows[:start]]
            vals = [r[j : j + 2] for r in rows[start:] if r[j] is not None or r[j + 1] is not None]
            dates = [a for a, b in vals if isinstance(a, datetime.datetime)]
            nums = [b for a, b in vals if isinstance(b, (int, float))]
            texts = collections.Counter(
                str(b) for a, b in vals if b is not None and not isinstance(b, (int, float))
            )
            print(
                "PAIR",
                j + 1,
                "HEAD",
                heads,
                "N",
                len(vals),
                "NUM",
                len(nums),
                "DATES",
                str(min(dates)) if dates else "",
                str(max(dates)) if dates else "",
                "MINMAX",
                (min(nums), max(nums)) if nums else "",
                "TEXT",
                texts.most_common(6),
            )
for p in (root / "data").glob("*.csv"):
    d = pd.read_csv(p, low_memory=False)
    t = pd.to_datetime(d["date"])
    x = d.drop(columns=[c for c in d if c == "date" or c.startswith("Unnamed:")])
    num = x.apply(pd.to_numeric, errors="coerce")
    print(
        "\nCSV",
        p.name,
        "SHAPE",
        d.shape,
        "TIME",
        str(t.min()),
        str(t.max()),
        "DUP",
        int(t.duplicated().sum()),
        "STEPS",
        t.diff().value_counts().head().to_dict(),
    )
    print("MISSING", x.isna().sum().to_dict())
    print(
        "NONNUM",
        {
            c: x.loc[num[c].isna() & x[c].notna(), c].value_counts().head().to_dict()
            for c in x
            if (num[c].isna() & x[c].notna()).any()
        },
    )
    print("QUANTILES", num.quantile([0, 0.5, 1]).round(3).to_dict(), flush=True)
