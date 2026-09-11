"""Prepare immutable source files into audited, versioned Parquet datasets."""

from __future__ import annotations

import hashlib
import json
import math
import re
from collections import Counter
from collections.abc import Callable, Iterable, Iterator, Mapping
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import numpy as np
import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import yaml

PREPARATION_CODE_VERSION = "data-preparation-v1"

TELEMETRY_COLUMNS = (
    "signal_id",
    "measured_at",
    "available_at",
    "value",
    "unit",
    "source",
    "quality",
    "issues",
    "source_column",
)
ANALYSIS_COLUMNS = (
    "signal_id",
    "point_id",
    "measured_at",
    "available_at",
    "value",
    "unit",
    "source",
    "quality",
    "issues",
    "source_label",
    "original_unit",
    "original_value",
)
EVENT_COLUMNS = (
    "signal_id",
    "point_id",
    "measured_at",
    "available_at",
    "source",
    "event_type",
    "raw_value",
    "source_label",
    "original_unit",
    "issues",
)

ANALYTE_NAMES = {
    "CFPP": "cfpp",
    "FilterabilityLimit.T": "cfpp",
    "90%.T": "distillation_t90",
    "50%.T": "distillation_t50",
    "95%.T": "distillation_t95",
    "98%.T": "distillation_t98",
    "EBP.T": "distillation_ebp",
    "IBP.T": "distillation_ibp",
    "CloudPoint": "cloud_point",
    "CloudPoint_1": "cloud_point_1",
    "PourPoint": "pour_point",
    "FlashPoint": "flash_point",
    "D15": "density_d15",
    "I250": "recovery_250",
    "I350": "recovery_350",
    "Mass.Sulfur": "sulfur_mass_fraction",
    "Mg.Sulfur": "sulfur",
    "CetaneNumber": "cetane_number",
}

EXPECTED_UNIT_KINDS = {
    "cfpp": "temperature",
    "distillation_t90": "temperature",
    "distillation_t50": "temperature",
    "distillation_t95": "temperature",
    "distillation_t98": "temperature",
    "distillation_ebp": "temperature",
    "distillation_ibp": "temperature",
    "cloud_point": "temperature",
    "cloud_point_1": "temperature",
    "pour_point": "temperature",
    "flash_point": "temperature",
    "density_d15": "density",
    "recovery_250": "volume_fraction",
    "recovery_350": "volume_fraction",
    "sulfur_mass_fraction": "mass_fraction",
    "sulfur": "mass_concentration",
    "cetane_number": "cetane_number",
}

UNIT_NORMALIZATION = {
    "°С": ("temperature", "°C", "identity"),
    "°C": ("temperature", "°C", "identity"),
    "кг/м3": ("density", "kg/m³", "identity"),
    "кг/м³": ("density", "kg/m³", "identity"),
    "% об.": ("volume_fraction", "% vol.", "identity"),
    "% масс.": ("mass_fraction", "% mass", "identity"),
    "мг/кг": ("mass_concentration", "mg/kg", "identity"),
    "ед.цет.ч.": ("cetane_number", "cetane number", "identity"),
    "ppm": ("mass_concentration", "mg/kg", "1 mass ppm = 1 mg/kg (assumption)"),
}

TELEMETRY_SCHEMA = pa.schema(
    [
        ("signal_id", pa.string()),
        ("measured_at", pa.timestamp("ns", tz="UTC")),
        ("available_at", pa.timestamp("ns", tz="UTC")),
        ("value", pa.float64()),
        ("unit", pa.string()),
        ("source", pa.string()),
        ("quality", pa.string()),
        ("issues", pa.string()),
        ("source_column", pa.string()),
    ]
)


@dataclass(frozen=True)
class PointMetadata:
    installation: str
    point: str
    product: str
    point_id: str


@dataclass(frozen=True)
class SeriesMetadata:
    signal_id: str
    point_id: str
    source_label: str
    original_unit: str
    canonical_unit: str | None
    conversion: str
    verification_status: str
    issues: tuple[str, ...]
    evidence: str


def _json_issues(issues: Iterable[str]) -> str:
    return json.dumps(list(issues), ensure_ascii=False, separators=(",", ":"))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _canonical_hash(value: Mapping[str, Any]) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _as_utc(value: Any, source_timezone: ZoneInfo) -> datetime:
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    if isinstance(value, date) and not isinstance(value, datetime):
        value = datetime.combine(value, datetime.min.time())
    if not isinstance(value, datetime):
        raise ValueError(f"unsupported timestamp {value!r}")
    if value.tzinfo is None or value.utcoffset() is None:
        value = value.replace(tzinfo=source_timezone)
    return value.astimezone(UTC)


def _slug(value: str) -> str:
    lowered = value.strip().lower().replace("ё", "е")
    lowered = lowered.replace("%", "_pct_")
    lowered = re.sub(r"[^a-zа-я0-9]+", "_", lowered, flags=re.IGNORECASE)
    return lowered.strip("_") or "unknown"


def _installation_namespace(value: str) -> str:
    if "Гидроочист" in value:
        return "ht"
    if "АВТ" in value:
        return "avt"
    return _slug(value)


def parse_point_header(value: str) -> PointMetadata:
    pattern = re.compile(
        r"Установка\s+'(?P<installation>[^']+)'\.*\s*"
        r"Точка отбора\s+'(?P<point>[^']+)'\.\s*"
        r"Продукт\s+'(?P<product>[^']+)'",
        re.IGNORECASE,
    )
    match = pattern.search(value)
    if match is None:
        raise ValueError(f"cannot parse LIMS point header: {value!r}")
    installation = match.group("installation")
    point = match.group("point")
    namespace = _installation_namespace(installation)
    return PointMetadata(
        installation=installation,
        point=point,
        product=match.group("product").rstrip("."),
        point_id=f"{namespace}:sample_point_{_slug(point)}",
    )


def _series_metadata(
    *,
    point: PointMetadata,
    source_label: str,
    original_unit: str,
    evidence: str,
) -> SeriesMetadata:
    analyte = ANALYTE_NAMES.get(source_label, _slug(source_label))
    namespace = point.point_id.split(":", maxsplit=1)[0]
    signal_id = f"lab:{namespace}.point_{_slug(point.point)}.{analyte}"
    expected_kind = EXPECTED_UNIT_KINDS.get(analyte)
    unit_info = UNIT_NORMALIZATION.get(original_unit)
    issues: list[str] = []
    canonical_unit: str | None = None
    conversion = "none"
    verification_status = "unverified"
    if unit_info is None:
        issues.append("unknown_unit")
    elif expected_kind is not None and unit_info[0] != expected_kind:
        issues.append("unit_mismatch")
        conversion = "blocked: source unit is incompatible with the indicator"
        verification_status = "suspect"
    else:
        canonical_unit = unit_info[1]
        conversion = unit_info[2]
        verification_status = "source_verified"
    return SeriesMetadata(
        signal_id=signal_id,
        point_id=point.point_id,
        source_label=source_label,
        original_unit=original_unit,
        canonical_unit=canonical_unit,
        conversion=conversion,
        verification_status=verification_status,
        issues=tuple(issues),
        evidence=evidence,
    )


def _value_quality(
    analyte: str,
    value: float,
    base_issues: Iterable[str],
    *,
    nonnegative_analytes: Iterable[str] = (
        "density_d15",
        "sulfur",
        "sulfur_mass_fraction",
        "cetane_number",
    ),
    percentage_analytes: Iterable[str] = ("recovery_250", "recovery_350"),
) -> tuple[str, str]:
    issues = list(base_issues)
    if analyte in nonnegative_analytes and value < 0:
        issues.append("negative_nonnegative_quantity")
    if analyte in percentage_analytes and not 0 <= value <= 100:
        issues.append("percentage_out_of_range")
    return ("suspect" if issues else "valid", _json_issues(issues))


def load_lims(
    path: Path,
    *,
    source_timezone: str,
    delay_minutes: int,
    quality_rules: Mapping[str, Any] | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, list[dict[str, Any]], list[dict[str, Any]]]:
    """Load every LIMS date/value pair independently."""

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    timezone = ZoneInfo(source_timezone)
    analyses: list[dict[str, Any]] = []
    events: list[dict[str, Any]] = []
    dictionary: list[dict[str, Any]] = []
    audit: list[dict[str, Any]] = []
    current_header: str | None = None
    effective_rules = quality_rules or {
        "nonnegative_analytes": (
            "density_d15",
            "sulfur",
            "sulfur_mass_fraction",
            "cetane_number",
        ),
        "percentage_analytes": ("recovery_250", "recovery_350"),
    }

    for date_index in range(0, len(rows[0]), 2):
        date_column = date_index + 1
        current_header = rows[0][date_index] or current_header
        if not current_header:
            raise ValueError(f"missing LIMS point header before column {date_column}")
        point = parse_point_header(str(current_header))
        source_label = str(rows[1][date_index] or "").strip()
        original_unit = str(rows[2][date_index] or "").strip()
        declared_count = rows[3][date_index + 1]
        evidence = f"{path.as_posix()}!{sheet.title}:{date_column}-{date_column + 1}"
        metadata = _series_metadata(
            point=point,
            source_label=source_label,
            original_unit=original_unit,
            evidence=evidence,
        )
        analyte = metadata.signal_id.rsplit(".", maxsplit=1)[-1]
        series_times: list[datetime] = []
        numeric_count = 0
        event_count = 0
        minimum: float | None = None
        maximum: float | None = None
        quality_counts: Counter[str] = Counter()

        dictionary.append(
            {
                "source": "lims",
                "source_column": f"{date_column}:{date_column + 1}",
                "signal_id": metadata.signal_id,
                "point_id": metadata.point_id,
                "description": source_label,
                "product": point.product,
                "original_unit": original_unit,
                "canonical_unit": metadata.canonical_unit,
                "conversion": metadata.conversion,
                "verification_status": metadata.verification_status,
                "evidence": evidence,
            }
        )

        for row in rows[4:]:
            raw_time = row[date_index]
            raw_value = row[date_index + 1]
            if raw_time is None and raw_value is None:
                continue
            try:
                measured_at = _as_utc(raw_time, timezone)
                available_at = measured_at + timedelta(minutes=delay_minutes)
                timestamp_issue: tuple[str, ...] = ()
                series_times.append(measured_at)
            except ValueError:
                measured_at = None
                available_at = None
                timestamp_issue = ("invalid_timestamp",)

            if isinstance(raw_value, (int, float)) and not isinstance(raw_value, bool):
                value = float(raw_value)
                if math.isfinite(value) and measured_at is not None:
                    quality, issues = _value_quality(
                        analyte,
                        value,
                        metadata.issues,
                        nonnegative_analytes=effective_rules["nonnegative_analytes"],
                        percentage_analytes=effective_rules["percentage_analytes"],
                    )
                    analyses.append(
                        {
                            "signal_id": metadata.signal_id,
                            "point_id": metadata.point_id,
                            "measured_at": measured_at,
                            "available_at": available_at,
                            "value": value,
                            "unit": metadata.canonical_unit,
                            "source": "lims",
                            "quality": quality,
                            "issues": issues,
                            "source_label": source_label,
                            "original_unit": original_unit,
                            "original_value": repr(raw_value),
                        }
                    )
                    numeric_count += 1
                    minimum = value if minimum is None else min(minimum, value)
                    maximum = value if maximum is None else max(maximum, value)
                    quality_counts[quality] += 1
                    continue

            event_issues = (*metadata.issues, *timestamp_issue)
            if raw_value is None:
                event_type = "missing_value"
                event_issues = (*event_issues, "missing_value")
            elif isinstance(raw_value, str):
                event_type = "source_text"
                event_issues = (*event_issues, "non_numeric_value")
            else:
                event_type = "invalid_value"
                event_issues = (*event_issues, "non_finite_or_invalid_value")
            events.append(
                {
                    "signal_id": metadata.signal_id,
                    "point_id": metadata.point_id,
                    "measured_at": measured_at,
                    "available_at": available_at,
                    "source": "lims",
                    "event_type": event_type,
                    "raw_value": None if raw_value is None else str(raw_value),
                    "source_label": source_label,
                    "original_unit": original_unit,
                    "issues": _json_issues(event_issues),
                }
            )
            event_count += 1

        duplicate_count = len(series_times) - len(set(series_times))
        audit.append(
            {
                "signal_id": metadata.signal_id,
                "declared_count": int(declared_count) if isinstance(declared_count, int) else None,
                "numeric_count": numeric_count,
                "event_count": event_count,
                "observed_count": numeric_count + event_count,
                "count_matches_declared": declared_count == numeric_count + event_count,
                "duplicate_timestamps": duplicate_count,
                "measured_at_min": min(series_times).isoformat() if series_times else None,
                "measured_at_max": max(series_times).isoformat() if series_times else None,
                "minimum": minimum,
                "maximum": maximum,
                "canonical_unit": metadata.canonical_unit,
                "quality_counts": dict(sorted(quality_counts.items())),
                "verification_status": metadata.verification_status,
            }
        )

    return (
        pd.DataFrame(analyses, columns=ANALYSIS_COLUMNS),
        pd.DataFrame(events, columns=EVENT_COLUMNS),
        dictionary,
        audit,
    )


def load_pak(
    path: Path,
    *,
    source_timezone: str,
    delay_minutes: int,
    quality_rules: Mapping[str, Any] | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, list[dict[str, Any]], list[dict[str, Any]]]:
    """Load the independent A:B and D:E PAK series around the empty C separator."""

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    timezone = ZoneInfo(source_timezone)
    definitions = {
        "24-2000:Mg.Sulfur": ("pak:ht.product_sulfur", "sulfur"),
        "24-2000:D15": ("pak:ht.product_density_d15", "density_d15"),
    }
    analyses: list[dict[str, Any]] = []
    events: list[dict[str, Any]] = []
    dictionary: list[dict[str, Any]] = []
    audit: list[dict[str, Any]] = []
    effective_rules = quality_rules or {
        "nonnegative_analytes": (
            "density_d15",
            "sulfur",
            "sulfur_mass_fraction",
            "cetane_number",
        ),
        "percentage_analytes": ("recovery_250", "recovery_350"),
    }

    if any(row[2] is not None for row in rows):
        raise ValueError("PAK separator column C is expected to be empty")

    for date_index in (0, 3):
        date_column = date_index + 1
        source_label = str(rows[0][date_index] or "").strip()
        if source_label not in definitions:
            raise ValueError(f"unsupported PAK series {source_label!r}")
        signal_id, analyte = definitions[source_label]
        original_unit = str(rows[1][date_index] or "").strip()
        unit_info = UNIT_NORMALIZATION.get(original_unit)
        expected_kind = EXPECTED_UNIT_KINDS[analyte]
        compatible = unit_info is not None and unit_info[0] == expected_kind
        canonical_unit = unit_info[1] if compatible else None
        conversion = unit_info[2] if compatible else "blocked: incompatible or unknown unit"
        verification_status = "source_verified" if compatible else "suspect"
        if compatible and "assumption" in conversion:
            verification_status = "source_verified_with_unit_assumption"
        evidence = f"{path.as_posix()}!{sheet.title}:{date_column}-{date_column + 1}"
        dictionary.append(
            {
                "source": "pak",
                "source_column": f"{date_column}:{date_column + 1}",
                "signal_id": signal_id,
                "point_id": "ht:product_outlet",
                "description": source_label,
                "product": "Гидроочищенное дизельное топливо",
                "original_unit": original_unit,
                "canonical_unit": canonical_unit,
                "conversion": conversion,
                "verification_status": verification_status,
                "evidence": evidence,
            }
        )
        series_times: list[datetime] = []
        numeric_count = 0
        event_count = 0
        minimum: float | None = None
        maximum: float | None = None
        quality_counts: Counter[str] = Counter()
        base_issues = () if compatible else ("unit_mismatch",)
        for row in rows[2:]:
            raw_time = row[date_index]
            raw_value = row[date_index + 1]
            if raw_time is None and raw_value is None:
                continue
            try:
                measured_at = _as_utc(raw_time, timezone)
                available_at = measured_at + timedelta(minutes=delay_minutes)
                series_times.append(measured_at)
            except ValueError:
                measured_at = None
                available_at = None
            if isinstance(raw_value, (int, float)) and not isinstance(raw_value, bool):
                value = float(raw_value)
                if math.isfinite(value) and measured_at is not None:
                    quality, issues = _value_quality(
                        analyte,
                        value,
                        base_issues,
                        nonnegative_analytes=effective_rules["nonnegative_analytes"],
                        percentage_analytes=effective_rules["percentage_analytes"],
                    )
                    analyses.append(
                        {
                            "signal_id": signal_id,
                            "point_id": "ht:product_outlet",
                            "measured_at": measured_at,
                            "available_at": available_at,
                            "value": value,
                            "unit": canonical_unit,
                            "source": "pak",
                            "quality": quality,
                            "issues": issues,
                            "source_label": source_label,
                            "original_unit": original_unit,
                            "original_value": repr(raw_value),
                        }
                    )
                    numeric_count += 1
                    minimum = value if minimum is None else min(minimum, value)
                    maximum = value if maximum is None else max(maximum, value)
                    quality_counts[quality] += 1
                    continue
            issue = "invalid_timestamp" if measured_at is None else "non_numeric_value"
            events.append(
                {
                    "signal_id": signal_id,
                    "point_id": "ht:product_outlet",
                    "measured_at": measured_at,
                    "available_at": available_at,
                    "source": "pak",
                    "event_type": "invalid_value",
                    "raw_value": None if raw_value is None else str(raw_value),
                    "source_label": source_label,
                    "original_unit": original_unit,
                    "issues": _json_issues((*base_issues, issue)),
                }
            )
            event_count += 1
        audit.append(
            {
                "signal_id": signal_id,
                "numeric_count": numeric_count,
                "event_count": event_count,
                "duplicate_timestamps": len(series_times) - len(set(series_times)),
                "measured_at_min": min(series_times).isoformat() if series_times else None,
                "measured_at_max": max(series_times).isoformat() if series_times else None,
                "minimum": minimum,
                "maximum": maximum,
                "canonical_unit": canonical_unit,
                "quality_counts": dict(sorted(quality_counts.items())),
            }
        )
    return (
        pd.DataFrame(analyses, columns=ANALYSIS_COLUMNS),
        pd.DataFrame(events, columns=EVENT_COLUMNS),
        dictionary,
        audit,
    )


def load_tag_dictionary(path: Path) -> tuple[dict[str, str], dict[str, str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["КИП"]
    avt: dict[str, str] = {}
    ht: dict[str, str] = {}
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 2).value:
            avt[str(sheet.cell(row, 2).value)] = str(sheet.cell(row, 1).value or "")
        if sheet.cell(row, 4).value:
            ht[str(sheet.cell(row, 4).value)] = str(sheet.cell(row, 3).value or "")
    return avt, ht


def _telemetry_dictionary_entries(
    *,
    source_path: Path,
    namespace: str,
    columns: Iterable[str],
    descriptions: Mapping[str, str],
) -> list[dict[str, Any]]:
    entries = []
    for column in columns:
        status = "suspect" if namespace == "ht" and column in {"T6", "P8"} else "unverified"
        evidence = f"context/Теги_хакатон.xlsx!КИП and {source_path.as_posix()}:{column}"
        entries.append(
            {
                "source": "telemetry",
                "source_column": column,
                "signal_id": f"{namespace}:{column}",
                "point_id": None,
                "description": descriptions.get(column),
                "product": None,
                "original_unit": None,
                "canonical_unit": None,
                "conversion": "none",
                "verification_status": status,
                "evidence": evidence,
            }
        )
    return entries


def _iter_telemetry_chunks(
    path: Path,
    *,
    namespace: str,
    descriptions: Mapping[str, str],
    source_timezone: str,
    delay_minutes: int,
    timestamp_column: str,
    timestamp_format: str,
    chunksize: int,
    service_pattern: re.Pattern[str],
    audit_exact_values: Iterable[float],
    negative_flow_enabled: bool = True,
    negative_flow_issue: str = "negative_flow",
) -> Iterator[tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]]:
    timezone = ZoneInfo(source_timezone)
    source = pd.read_csv(path, low_memory=False)
    service_columns = [column for column in source if service_pattern.search(str(column))]
    source = source.drop(columns=service_columns)
    if timestamp_column not in source:
        raise ValueError(f"{path}: missing timestamp column {timestamp_column!r}")
    parsed = pd.to_datetime(source[timestamp_column], format=timestamp_format, errors="coerce")
    invalid_timestamp_count = int(parsed.isna().sum())
    if invalid_timestamp_count:
        raise ValueError(f"{path}: {invalid_timestamp_count} invalid timestamps")
    source = source.assign(_parsed_timestamp=parsed).sort_values("_parsed_timestamp", kind="stable")
    for start in range(0, len(source), chunksize):
        chunk = source.iloc[start : start + chunksize].copy()
        parsed_chunk = chunk.pop("_parsed_timestamp")
        chunk = chunk.drop(columns=[timestamp_column])
        measured_at = parsed_chunk.dt.tz_localize(timezone).dt.tz_convert("UTC")
        signal_columns = list(chunk.columns)
        long = chunk.assign(measured_at=measured_at).melt(
            id_vars="measured_at", var_name="source_column", value_name="raw_value"
        )
        long["value"] = pd.to_numeric(long["raw_value"], errors="coerce")
        nonnumeric = long["value"].isna() & long["raw_value"].notna()
        nonfinite = long["value"].notna() & ~np.isfinite(long["value"])
        missing = long["raw_value"].isna()
        invalid = nonnumeric | nonfinite
        event_rows = long.loc[invalid, ["measured_at", "source_column", "raw_value"]].copy()
        if not event_rows.empty:
            event_rows["signal_id"] = namespace + ":" + event_rows["source_column"].astype(str)
            event_rows["point_id"] = None
            event_rows["available_at"] = event_rows["measured_at"] + pd.Timedelta(
                minutes=delay_minutes
            )
            event_rows["source"] = "telemetry"
            event_rows["event_type"] = "invalid_value"
            event_rows["raw_value"] = event_rows["raw_value"].astype(str)
            event_rows["source_label"] = event_rows["source_column"]
            event_rows["original_unit"] = None
            event_rows["issues"] = _json_issues(("non_numeric_or_non_finite_value",))
            event_rows = event_rows.loc[:, list(EVENT_COLUMNS)]

        long["signal_id"] = namespace + ":" + long["source_column"].astype(str)
        long["available_at"] = long["measured_at"] + pd.Timedelta(minutes=delay_minutes)
        long["unit"] = None
        long["source"] = "telemetry"
        long["quality"] = "valid"
        long["issues"] = "[]"
        long.loc[missing, "quality"] = "missing"
        long.loc[missing, "issues"] = _json_issues(("missing_value",))
        flow_columns = (
            {column for column, description in descriptions.items() if "Расход" in description}
            if negative_flow_enabled
            else set()
        )
        negative_flow = long["source_column"].isin(flow_columns) & (long["value"] < 0)
        long.loc[negative_flow, "quality"] = "suspect"
        long.loc[negative_flow, "issues"] = _json_issues((negative_flow_issue,))
        long = long.loc[~invalid, list(TELEMETRY_COLUMNS)]

        per_signal: dict[str, Any] = {}
        for column in signal_columns:
            values = pd.to_numeric(chunk[column], errors="coerce")
            latest_raw = chunk[column].iloc[-1]
            latest_numeric = values.iloc[-1]
            if pd.isna(latest_raw):
                latest_value = None
                latest_quality = "missing"
                latest_issues = _json_issues(("missing_value",))
            elif pd.isna(latest_numeric) or not math.isfinite(float(latest_numeric)):
                latest_value = None
                latest_quality = "missing"
                latest_issues = _json_issues(("non_numeric_or_non_finite_value",))
            else:
                latest_value = float(latest_numeric)
                if column in flow_columns and latest_value < 0:
                    latest_quality = "suspect"
                    latest_issues = _json_issues((negative_flow_issue,))
                else:
                    latest_quality = "valid"
                    latest_issues = "[]"
            exact_counts = {
                str(exact): int((values == exact).sum()) for exact in audit_exact_values
            }
            per_signal[f"{namespace}:{column}"] = {
                "missing_count": int(values.isna().sum()),
                "minimum": float(values.min()) if values.notna().any() else None,
                "maximum": float(values.max()) if values.notna().any() else None,
                "negative_count": int((values < 0).sum()),
                "exact_value_counts": exact_counts,
                "suspect_count": int((values < 0).sum()) if column in flow_columns else 0,
                "latest_value": latest_value,
                "latest_quality": latest_quality,
                "latest_issues": latest_issues,
            }
        chunk_audit = {
            "input_rows": len(chunk),
            "service_columns_removed": service_columns,
            "measured_at_min": measured_at.min().isoformat(),
            "measured_at_max": measured_at.max().isoformat(),
            "timestamps": measured_at,
            "per_signal": per_signal,
        }
        yield long, event_rows, chunk_audit


def _merge_signal_audit(target: dict[str, Any], update: Mapping[str, Any]) -> None:
    for signal_id, metrics in update.items():
        current = target.setdefault(
            signal_id,
            {
                "missing_count": 0,
                "minimum": None,
                "maximum": None,
                "negative_count": 0,
                "exact_value_counts": {},
                "suspect_count": 0,
                "latest_value": None,
                "latest_quality": "missing",
                "latest_issues": _json_issues(("missing_value",)),
            },
        )
        current["missing_count"] += metrics["missing_count"]
        current["negative_count"] += metrics["negative_count"]
        current["suspect_count"] += metrics["suspect_count"]
        current["latest_value"] = metrics["latest_value"]
        current["latest_quality"] = metrics["latest_quality"]
        current["latest_issues"] = metrics["latest_issues"]
        if metrics["minimum"] is not None:
            current["minimum"] = (
                metrics["minimum"]
                if current["minimum"] is None
                else min(current["minimum"], metrics["minimum"])
            )
            current["maximum"] = (
                metrics["maximum"]
                if current["maximum"] is None
                else max(current["maximum"], metrics["maximum"])
            )
        for value, count in metrics["exact_value_counts"].items():
            current["exact_value_counts"][value] = (
                current["exact_value_counts"].get(value, 0) + count
            )


def _write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def prepare_data(
    root: Path | None = None, *, progress: Callable[[str], None] | None = None
) -> dict[str, Any]:
    """Run the complete point-B preparation using repository-owned paths."""

    repository = (root or Path(__file__).resolve().parents[3]).resolve()
    data_config_path = repository / "config" / "data.yaml"
    model_config_path = repository / "config" / "model.yaml"
    data_config = yaml.safe_load(data_config_path.read_text(encoding="utf-8"))
    model_config = yaml.safe_load(model_config_path.read_text(encoding="utf-8"))
    output_dir = repository / data_config["output_directory"]
    output_dir.mkdir(parents=True, exist_ok=True)

    source_paths = {
        "lims": repository / data_config["sources"]["lims"],
        "pak": repository / data_config["sources"]["pak"],
        "tag_dictionary": repository / data_config["sources"]["tag_dictionary"],
    }
    for item in data_config["sources"]["telemetry"]:
        source_paths[f"telemetry_{item['namespace']}"] = repository / item["path"]
    input_hashes = {
        name: {
            "path": path.relative_to(repository).as_posix(),
            "sha256": _sha256(path),
            "bytes": path.stat().st_size,
        }
        for name, path in sorted(source_paths.items())
    }
    version_inputs = {
        "preparation_code_version": PREPARATION_CODE_VERSION,
        "preparation_code_sha256": _sha256(Path(__file__)),
        "input_hashes": input_hashes,
        "data_config": data_config,
        "model_time_and_availability": {
            "time": model_config["time"],
            "availability": model_config["availability"],
        },
    }
    dataset_version = f"sha256:{_canonical_hash(version_inputs)}"

    source_timezone = model_config["time"]["source_timezone"]
    delays = model_config["availability"]["default_delay_minutes"]
    descriptions_by_namespace = dict(
        zip(("avt", "ht"), load_tag_dictionary(source_paths["tag_dictionary"]), strict=True)
    )
    dictionary: list[dict[str, Any]] = []
    events_frames: list[pd.DataFrame] = []
    telemetry_audit: list[dict[str, Any]] = []
    telemetry_path = output_dir / "telemetry.parquet"
    writer = pq.ParquetWriter(
        telemetry_path,
        TELEMETRY_SCHEMA,
        compression=data_config["parquet"]["compression"],
    )
    telemetry_rows = 0
    telemetry_max: datetime | None = None
    try:
        for source in data_config["sources"]["telemetry"]:
            source_path = repository / source["path"]
            if progress:
                progress(f"Preparing telemetry source {source['path']}")
            header = pd.read_csv(source_path, nrows=0).columns.tolist()
            pattern = re.compile(data_config["csv"]["service_column_pattern"])
            signal_columns = [
                column
                for column in header
                if column != data_config["csv"]["timestamp_column"] and not pattern.search(column)
            ]
            descriptions = descriptions_by_namespace[source["namespace"]]
            dictionary.extend(
                _telemetry_dictionary_entries(
                    source_path=Path(source["path"]),
                    namespace=source["namespace"],
                    columns=signal_columns,
                    descriptions=descriptions,
                )
            )
            summary: dict[str, Any] = {
                "path": source["path"],
                "namespace": source["namespace"],
                "input_rows": 0,
                "signal_count": len(signal_columns),
                "service_columns_removed": [],
                "duplicate_timestamps": 0,
                "unexpected_step_count": 0,
                "step_counts_seconds": {},
                "measured_at_min": None,
                "measured_at_max": None,
                "per_signal": {},
            }
            previous_time: pd.Timestamp | None = None
            for frame, event_frame, chunk_audit in _iter_telemetry_chunks(
                source_path,
                namespace=source["namespace"],
                descriptions=descriptions,
                source_timezone=source_timezone,
                delay_minutes=delays["online"],
                timestamp_column=data_config["csv"]["timestamp_column"],
                timestamp_format=data_config["csv"]["timestamp_format"],
                chunksize=data_config["csv"]["chunksize"],
                service_pattern=pattern,
                audit_exact_values=data_config["quality_rules"]["audit_exact_values"],
                negative_flow_enabled=data_config["quality_rules"]["negative_flow"]["enabled"],
                negative_flow_issue=data_config["quality_rules"]["negative_flow"]["issue"],
            ):
                table = pa.Table.from_pandas(frame, schema=TELEMETRY_SCHEMA, preserve_index=False)
                writer.write_table(table)
                telemetry_rows += len(frame)
                if not event_frame.empty:
                    events_frames.append(event_frame)
                times = chunk_audit.pop("timestamps")
                all_times = (
                    times
                    if previous_time is None
                    else pd.concat([pd.Series([previous_time]), times], ignore_index=True)
                )
                deltas = all_times.diff().dropna().dt.total_seconds().astype(int)
                counts = Counter(deltas.tolist())
                for seconds, count in counts.items():
                    key = str(seconds)
                    summary["step_counts_seconds"][key] = (
                        summary["step_counts_seconds"].get(key, 0) + count
                    )
                expected_seconds = data_config["csv"]["expected_step_minutes"] * 60
                summary["duplicate_timestamps"] += counts.get(0, 0)
                summary["unexpected_step_count"] += sum(
                    count for seconds, count in counts.items() if seconds != expected_seconds
                )
                previous_time = times.iloc[-1]
                summary["input_rows"] += chunk_audit["input_rows"]
                summary["service_columns_removed"] = sorted(
                    set(summary["service_columns_removed"])
                    | set(chunk_audit["service_columns_removed"])
                )
                summary["measured_at_min"] = (
                    summary["measured_at_min"] or chunk_audit["measured_at_min"]
                )
                summary["measured_at_max"] = chunk_audit["measured_at_max"]
                _merge_signal_audit(summary["per_signal"], chunk_audit["per_signal"])
            telemetry_audit.append(summary)
            source_max = datetime.fromisoformat(summary["measured_at_max"])
            telemetry_max = source_max if telemetry_max is None else max(telemetry_max, source_max)
            if progress:
                progress(
                    f"Prepared {source['namespace']} telemetry: "
                    f"{summary['input_rows']} timestamps, {summary['signal_count']} signals"
                )
    finally:
        writer.close()

    if progress:
        progress("Preparing independent LIMS series")
    lims, lims_events, lims_dictionary, lims_audit = load_lims(
        source_paths["lims"],
        source_timezone=source_timezone,
        delay_minutes=delays["lims"],
        quality_rules=data_config["quality_rules"],
    )
    if progress:
        progress("Preparing independent PAK series")
    pak, pak_events, pak_dictionary, pak_audit = load_pak(
        source_paths["pak"],
        source_timezone=source_timezone,
        delay_minutes=delays["pak"],
        quality_rules=data_config["quality_rules"],
    )
    analyses = pd.concat([lims, pak], ignore_index=True).sort_values(
        ["measured_at", "signal_id"], kind="stable"
    )
    analyses.to_parquet(
        output_dir / "analyses.parquet",
        index=False,
        compression=data_config["parquet"]["compression"],
    )
    events_frames.extend([lims_events, pak_events])
    events = (
        pd.concat(events_frames, ignore_index=True)
        if events_frames
        else pd.DataFrame(columns=EVENT_COLUMNS)
    )
    if not events.empty:
        events = events.sort_values(["measured_at", "signal_id"], kind="stable", na_position="last")
    events.to_parquet(
        output_dir / "events.parquet",
        index=False,
        compression=data_config["parquet"]["compression"],
    )
    dictionary.extend(lims_dictionary)
    dictionary.extend(pak_dictionary)
    dictionary.sort(key=lambda item: (item["source"], item["signal_id"]))
    dictionary_payload = {
        "schema_version": data_config["schema_version"],
        "dataset_version": dataset_version,
        "fields": dictionary,
    }
    _write_json(output_dir / "data_dictionary.json", dictionary_payload)

    candidate_signals = [
        {
            "signal_id": item["signal_id"],
            "source": item["source"],
            "verification_status": item["verification_status"],
            "candidate_status": (
                "eligible_for_review"
                if item["verification_status"].startswith("source_verified")
                else "requires_verification"
            ),
            "reason": (
                "Source label and unit are compatible; model relevance is not yet established."
                if item["verification_status"].startswith("source_verified")
                else "Meaning or unit must be verified before physical calculations."
            ),
        }
        for item in dictionary
    ]
    _write_json(
        output_dir / "candidate_signals.json",
        {
            "dataset_version": dataset_version,
            "signals": candidate_signals,
            "control_availability_declared": False,
        },
    )

    state_values: dict[str, dict[str, Any]] = {}
    for telemetry_source in telemetry_audit:
        for signal_id, signal_audit in telemetry_source["per_signal"].items():
            state_values[signal_id] = {
                "value": signal_audit["latest_value"],
                "unit": None,
                "source": "telemetry",
                "measured_at": telemetry_source["measured_at_max"],
                "available_at": telemetry_source["measured_at_max"],
                "quality": signal_audit["latest_quality"],
                "issues": json.loads(signal_audit["latest_issues"]),
            }
    if telemetry_max is not None:
        available_analyses = analyses.loc[analyses["available_at"] <= telemetry_max]
        latest_analyses = (
            available_analyses.sort_values("available_at", kind="stable")
            .groupby("signal_id", sort=False)
            .tail(1)
        )
        for row in latest_analyses.itertuples(index=False):
            state_values[row.signal_id] = {
                "value": row.value,
                "unit": row.unit,
                "source": row.source,
                "measured_at": row.measured_at.isoformat(),
                "available_at": row.available_at.isoformat(),
                "quality": row.quality,
                "issues": json.loads(row.issues),
                "point_id": row.point_id,
            }
    _write_json(
        output_dir / "prepared_state_example.json",
        {
            "schema_version": data_config["schema_version"],
            "kind": "prepared_telemetry_state",
            "synthetic": False,
            "snapshot_contract_ready": False,
            "note": "Prepared real state for handoff; ProcessSnapshot is implemented in route C.",
            "mode": "replay",
            "as_of": telemetry_max.isoformat() if telemetry_max else None,
            "dataset_version": dataset_version,
            "values": dict(sorted(state_values.items())),
        },
    )

    sensitivity_minutes = model_config["availability"]["sensitivity_delay_minutes"]["lims"]
    sensitivity_cutoff = telemetry_max
    default_available = (
        int((lims["available_at"] <= sensitivity_cutoff).sum()) if sensitivity_cutoff else None
    )
    shifted_available = (
        int(
            (
                lims["measured_at"] + pd.Timedelta(minutes=sensitivity_minutes)
                <= sensitivity_cutoff
            ).sum()
        )
        if sensitivity_cutoff
        else None
    )
    audit_report = {
        "schema_version": data_config["schema_version"],
        "dataset_version": dataset_version,
        "preparation_code_version": PREPARATION_CODE_VERSION,
        "preparation_code_sha256": version_inputs["preparation_code_sha256"],
        "input_hashes": input_hashes,
        "assumptions": {
            "source_timezone": source_timezone,
            "storage_timezone": model_config["time"]["storage_timezone"],
            "default_delay_minutes": delays,
            "lims_delay_sensitivity_minutes": sensitivity_minutes,
            "exact_307_policy": "counted_for_audit_only; not automatically invalid",
        },
        "outputs": {
            "telemetry_rows": telemetry_rows,
            "analysis_rows": len(analyses),
            "event_rows": len(events),
            "dictionary_entries": len(dictionary),
        },
        "telemetry": telemetry_audit,
        "lims": lims_audit,
        "pak": pak_audit,
        "lims_delay_sensitivity": {
            "cutoff": sensitivity_cutoff.isoformat() if sensitivity_cutoff else None,
            "available_with_default_delay": default_available,
            "available_with_sensitivity_delay": shifted_available,
            "difference": (
                default_available - shifted_available
                if default_available is not None and shifted_available is not None
                else None
            ),
            "records_with_shifted_available_at": len(lims),
            "availability_shift_minutes": sensitivity_minutes,
            "interpretation": "Experimental availability shift, not a confirmed laboratory delay.",
        },
    }
    _write_json(output_dir / "audit_report.json", audit_report)
    return audit_report
